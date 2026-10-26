"""The Excel workbook the app offers for download (user, 2026-10-26: "może to wszystko do tego Excela
przenosić ... poukładane i edytowalne").

Built in the style of the user's own experiment sheet (Arial, yellow input cells, dark header row, PLN
formats): every proposed bet builder with its chance, fair price, "graj od" price, the app's settlement
and the paper result at "graj od", plus yellow columns for the user's own price, stake and status with
the same formulas as the user's sheet (payout after the 12% tax, profit, break-even chance, edge, EV).
Summary formulas at the top, a binomial "Symulacja" sheet like the user's, and a short how-to.  Formulas
are written as formulas (Excel computes them on opening: fullCalcOnLoad).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="17324D")
HEAD_FONT = Font(name=ARIAL, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BLOCK_FILL = PatternFill("solid", fgColor="EAF2F8")
PLN = '#,##0.00\\ "PLN";[Red]\\(#,##0.00\\ "PLN"\\)'
PCT = "0.0%"
DEC = "0.00"
FIRST_ROW = 10          # first data row of the "Zestawy" table
LAST_ROW = 5000         # formulas cover this many rows, so the user can append
WARSAW = dt.timezone(dt.timedelta(hours=1))


def _font(bold: bool = False, size: int = 10, color: str | None = None) -> Font:
    return Font(name=ARIAL, bold=bold, size=size, color=color)


def _local(iso: str) -> dt.datetime:
    return dt.datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(WARSAW)


def _leg_text(leg: dict) -> str:
    s = leg.get("opis", "")
    i = s.find(": ")
    return s[i + 2:] if i >= 0 else s


def _came(leg: dict) -> str:
    v = leg.get("wartosc")
    if v is None:
        return ""
    if leg.get("side") == "more":
        return f"{_leg_text(leg)}: {int(v)}" + (" goli" if leg.get("stat") == "goals" else "")
    if isinstance(v, str):
        return f"{_leg_text(leg)}: {v}"
    need = int(round(leg["line"] + 0.5)) if leg.get("line") is not None else None
    return f"{_leg_text(leg)}: {int(v)}" + (f" (potrzeba {need})" if need else "")


def write_workbook(path: Path, picks_out: dict, made: str) -> int:
    rows = []
    for day in sorted(picks_out):
        for b in picks_out[day].get("bb", []):
            rows.append(b)
    rows.sort(key=lambda b: (b["start"], b["mecz"], b.get("rodzaj", "")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Zestawy"
    ws["A1"] = "Typometr: bet buildery"
    ws["A1"].font = _font(True, 14)
    ws["A2"] = (f"Stan z {_local(made).strftime('%d.%m.%Y %H:%M')}. Żółte pola uzupełniasz ręcznie (kurs z Superbet, stawka, "
                "status); pozostałe kolumny liczą się same. Nowe zestawy dopisuj poniżej albo pobierz świeży arkusz z aplikacji.")
    ws["A2"].font = _font()
    for addr, text in (("A4", "Założenia"), ("F4", "Zestawy z aplikacji"), ("I4", "Twoje zakłady")):
        ws[addr] = text
        ws[addr].font = _font(True)
        ws[addr].fill = BLOCK_FILL
    ws["A5"], ws["C5"] = "Podatek od stawki", 0.12
    ws["A6"], ws["C6"] = "Stawka na papierze (zł)", 100
    for addr, fmt in (("C5", PCT), ("C6", '#,##0.00\\ "PLN"')):
        ws[addr].fill = INPUT_FILL
        ws[addr].number_format = fmt
    K, M, N, O, P, R = (f"{c}{FIRST_ROW}:{c}{LAST_ROW}" for c in ("K", "M", "N", "O", "P", "R"))
    summary = [
        ("F5", "Rozliczone", "G5", f'=COUNTIF({K},"Weszło")+COUNTIF({K},"Nie weszło")', None),
        ("F6", "Weszło", "G6", f'=COUNTIF({K},"Weszło")', None),
        ("F7", "Skuteczność", "G7", '=IFERROR(G6/G5,"")', PCT),
        ("F8", "Zysk przy kursie graj od", "G8", f"=SUM({M})", PLN),
        ("I5", "Rozliczone", "J5", f'=COUNTIF({P},"Wygrany")+COUNTIF({P},"Przegrany")+COUNTIF({P},"Zwrot")', None),
        ("I6", "Postawione", "J6", f'=SUMIF({P},"Wygrany",{O})+SUMIF({P},"Przegrany",{O})+SUMIF({P},"Zwrot",{O})', PLN),
        ("I7", "Zysk / strata", "J7", f"=SUM({R})", PLN),
        ("I8", "ROI", "J8", '=IFERROR(J7/J6,"")', PCT),
    ]
    for la, label, va, formula, fmt in summary:
        ws[la] = label
        ws[la].font = _font()
        ws[va] = formula
        ws[va].font = _font(True)
        if fmt:
            ws[va].number_format = fmt
    headers = ["ID", "Data", "Godzina", "Mecz", "Drużyna", "Zdarzenia", "Rodzaj", "Szansa (model)", "Uczciwy kurs", "Graj od",
               "Wynik wg aplikacji", "Co padło", 'Zysk przy "graj od"', "Twój kurs", "Stawka", "Status", "Zwrot", "Zysk / strata",
               "Próg BE", "Przewaga", "EV"]
    hr = FIRST_ROW - 1
    for i, h in enumerate(headers, 1):
        c = ws.cell(hr, i, h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in range(FIRST_ROW, LAST_ROW + 1):
        ws[f"M{r}"] = f'=IF(K{r}="Weszło",$C$6*((1-$C$5)*J{r}-1),IF(K{r}="Nie weszło",-$C$6,""))'
        ws[f"Q{r}"] = f'=IF(OR(P{r}="",P{r}="Oczekuje",O{r}=""),"",IF(P{r}="Wygrany",O{r}*(1-$C$5)*N{r},IF(P{r}="Zwrot",O{r},0)))'
        ws[f"R{r}"] = f'=IF(Q{r}="","",Q{r}-O{r})'
        ws[f"S{r}"] = f'=IFERROR(1/((1-$C$5)*N{r}),"")'
        ws[f"T{r}"] = f'=IF(OR(N{r}="",H{r}=""),"",H{r}-S{r})'
        ws[f"U{r}"] = f'=IF(OR(N{r}="",O{r}=""),"",O{r}*(H{r}*(1-$C$5)*N{r}-1))'
        for col, fmt in (("M", PLN), ("Q", PLN), ("R", PLN), ("U", PLN), ("S", PCT), ("T", PCT), ("N", DEC), ("O", '#,##0.00\\ "PLN"')):
            ws[f"{col}{r}"].number_format = fmt
        for col in ("N", "O", "P"):
            ws[f"{col}{r}"].fill = INPUT_FILL
        for col in "ABCDEFGHIJKLMNOPQRSTU":
            ws[f"{col}{r}"].font = _font()
    for i, b in enumerate(rows, 1):
        r = FIRST_ROW + i - 1
        start = _local(b["start"])
        won = b.get("wynik")
        values = [i, start.date(), start.strftime("%H:%M"), b["mecz"], b["druzyna"], " + ".join(_leg_text(l) for l in b["legs"]),
                  "z wynikiem meczu" if b.get("rodzaj") == "z_wynikiem" else "bezpieczniejszy", b["p_model"], b.get("kurs_uczciwy"),
                  b.get("kurs_min_z_podatkiem"), "Weszło" if won == "wygrana" else "Nie weszło" if won == "przegrana" else "",
                  " · ".join(x for x in (_came(l) for l in b["legs"]) if x)]
        for j, v in enumerate(values, 1):
            ws.cell(r, j, v)
        ws[f"B{r}"].number_format = "DD.MM.YYYY"
        ws[f"H{r}"].number_format = PCT
        ws[f"I{r}"].number_format = DEC
        ws[f"J{r}"].number_format = DEC
    dv_status = DataValidation(type="list", formula1='"Oczekuje,Wygrany,Przegrany,Zwrot"', allow_blank=True)
    dv_odds = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="1.01", allow_blank=True)
    dv_stake = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0.01", allow_blank=True)
    for dv, col in ((dv_status, "P"), (dv_odds, "N"), (dv_stake, "O")):
        dv.add(f"{col}{FIRST_ROW}:{col}{LAST_ROW}")
        ws.add_data_validation(dv)
    widths = {"A": 5, "B": 11, "C": 8, "D": 30, "E": 20, "F": 60, "G": 16, "H": 11, "I": 10, "J": 9, "K": 13, "L": 50, "M": 15,
              "N": 10, "O": 11, "P": 11, "Q": 13, "R": 13, "S": 10, "T": 10, "U": 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"D{FIRST_ROW}"
    ws.auto_filter.ref = f"A{hr}:U{max(FIRST_ROW, FIRST_ROW + len(rows) - 1)}"

    # ---- Symulacja: the user's binomial sheet, pointed at this strategy
    sim = wb.create_sheet("Symulacja")
    sim["A1"] = "Symulacja strategii"
    sim["A1"].font = _font(True, 14)
    sim["A2"] = "Ile daje stała stawka na N zestawów przy danym kursie i skuteczności. Zmień żółte założenia."
    sim["A2"].font = _font()
    assumptions = [("Budżet", "B4", 1000, '#,##0.00\\ "PLN"', True), ("Liczba zestawów", "B5", 10, "0", True),
                   ("Stawka na zestaw", "B6", "=B4/B5", '#,##0.00\\ "PLN"', False), ("Średni kurs", "B7", 1.55, DEC, True),
                   ("Podatek od stawki", "B8", "=Zestawy!C5", PCT, False), ("Skuteczność (szacunek)", "B9", 0.76, PCT, True)]
    for i, (label, addr, val, fmt, inp) in enumerate(assumptions):
        sim[f"A{4 + i}"] = label
        sim[f"A{4 + i}"].font = _font()
        sim[addr] = val
        sim[addr].number_format = fmt
        sim[addr].font = _font()
        if inp:
            sim[addr].fill = INPUT_FILL
    outputs = [("Próg skuteczności (zero po podatku)", "=1/((1-B8)*B7)", PCT), ("Zysk na wygranym zestawie", "=B6*((1-B8)*B7-1)", PLN),
               ("Strata na przegranym", "=-B6", PLN), ("Minimalna liczba wygranych, żeby nie stracić", "=ROUNDUP(B5/((1-B8)*B7),0)", "0"),
               ("Oczekiwany wynik", "=B4*(B9*(1-B8)*B7-1)", PLN), ("Oczekiwane ROI", "=F8/B4", PCT),
               ("Odchylenie standardowe wyniku", "=B6*(1-B8)*B7*SQRT(B5*B9*(1-B9))", PLN), ("Szansa zakończenia na plusie", '=SUMIF(C15:C65,">0",B15:B65)', PCT)]
    for i, (label, formula, fmt) in enumerate(outputs):
        sim[f"E{4 + i}"] = label
        sim[f"E{4 + i}"].font = _font()
        sim[f"F{4 + i}"] = formula
        sim[f"F{4 + i}"].number_format = fmt
        sim[f"F{4 + i}"].font = _font(True)
    for j, h in enumerate(("Wygrane zestawy", "Prawdopodobieństwo", "Wynik całkowity"), 1):
        c = sim.cell(14, j, h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for k in range(0, 51):
        r = 15 + k
        sim[f"A{r}"] = k
        sim[f"B{r}"] = f'=IF(A{r}<=$B$5,COMBIN($B$5,A{r})*$B$9^A{r}*(1-$B$9)^($B$5-A{r}),"")'
        sim[f"C{r}"] = f'=IF(A{r}<=$B$5,A{r}*$B$6*(1-$B$8)*$B$7-$B$4,"")'
        sim[f"B{r}"].number_format = "0.00%"
        sim[f"C{r}"].number_format = PLN
        for col in "ABC":
            sim[f"{col}{r}"].font = _font()
    for col, w in (("A", 34), ("B", 18), ("C", 18), ("E", 40), ("F", 16)):
        sim.column_dimensions[col].width = w

    # ---- Jak korzystać
    how = wb.create_sheet("Jak korzystać")
    lines = [
        "Jak korzystać z arkusza",
        "",
        "1. Arkusz „Zestawy” to wszystkie bet buildery zaproponowane przez Typometr, od najstarszego. Szansa, uczciwy kurs i „graj od” są z modelu;",
        "   „Wynik wg aplikacji” i „Co padło” wypełniają się po meczu (w świeżo pobranym arkuszu).",
        "2. Gdy zagrasz zestaw w Superbet, wpisz w jego wierszu: Twój kurs, Stawka, Status (Oczekuje, a po meczu Wygrany / Przegrany / Zwrot).",
        "   Zwrot, Zysk, Próg BE, Przewaga i EV policzą się same, tak jak w Twoim arkuszu eksperymentu.",
        "3. Przewaga = szansa z modelu minus próg BE przy Twoim kursie. Dodatnia znaczy, że kurs jest co najmniej „graj od”.",
        "4. Podsumowanie u góry: ile zestawów weszło i jaki byłby wynik przy kursie „graj od” (100 zł każdy), a obok wynik Twoich zakładów.",
        "5. „Symulacja” pokazuje, czego spodziewać się po N zestawach przy stałej stawce: próg skuteczności, oczekiwany wynik i rozrzut.",
        "6. Wpisy zrobione w aplikacji (przycisk „Dodaj do zakładów”) nie przenoszą się do tego pliku automatycznie: to jest kopia do własnej analizy.",
        "   Najprościej: zakłady zapisuj w aplikacji (rozlicza je sama), a arkusz pobieraj, gdy chcesz mieć wszystko w Excelu.",
        "",
        f"Stan z {_local(made).strftime('%d.%m.%Y %H:%M')}. Podatek 12% od stawki (ustawa o grach hazardowych, art. 73–74).",
    ]
    for i, line in enumerate(lines, 1):
        how[f"A{i}"] = line
        how[f"A{i}"].font = _font(i == 1, 14 if i == 1 else 10)
    how.column_dimensions["A"].width = 130
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return len(rows)
