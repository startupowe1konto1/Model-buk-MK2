"""Polish bet ledger in Excel: singles and 2-3 leg coupons, Polish betting taxes, real and paper bets kept apart.

Writes ``Zaklady MK2.xlsx`` (default folder ``Documents\\Model-Buk-Data-MK2`` or ``$MODEL_BUK_MK2_OUT``) with two
sheets:

``Zakłady``
    An Excel table (header row frozen, filters on).  The user types only the inputs (Data, Godzina, Rozgrywki,
    Mecz, Rodzaj, Zakład, Bukmacher, Kurs, Stawka, Wynik, Szansa wg modelu, Źródło, Tryb, Notatki).  Every
    derived column (Lp, Podatek 12%, Stawka po podatku, Wypłata, Podatek od wygranej, Zysk, Saldo, Min. kurs wg
    modelu) is a formula pre-filled for ``--rows`` rows; it stays blank while a row is empty or not settled.
``Podsumowanie``
    SUMIFS/COUNTIFS summaries, separate for Tryb = Realny and Tryb = Papier: totals, by Rodzaj, by Źródło and by
    month (12 months from ``--first-month``), with a short "Jak używać" note on top.

Polish rules written into the formulas:
    * 12% of the stake is taken at placement, so the effective stake is 0.88 x stake;
    * a win (W) pays 0.88 x stake x odds, a loss (P) pays 0, a void (Z, zwrot) returns the full stake;
    * a winning payout above 2280 zl pays 10% income tax on the whole payout (column "Podatek od wygranej");
    * Zysk = Wypłata - Podatek od wygranej - Stawka; Saldo = running total of Zysk of settled bets, per Tryb;
    * Min. kurs wg modelu = 1 / (szansa x 0.88).

The workbook is the user's own ledger, so an existing file is never overwritten unless ``--force`` is given, and
even then the old file is first copied to ``<name>.backup-<timestamp>.xlsx`` next to it.

Usage (from app/):

    .venv\\Scripts\\python scripts\\mk2\\make_ledger.py
    .venv\\Scripts\\python scripts\\mk2\\make_ledger.py --out C:\\temp\\test.xlsx --rows 1500 --no-seed
    .venv\\Scripts\\python scripts\\mk2\\make_ledger.py --force     # rebuild from scratch, old file kept as backup
"""

from __future__ import annotations

import argparse
import io
import os
import shutil
import sys
from datetime import date, datetime, time
from pathlib import Path
from typing import NamedTuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_OUT_ROOT = Path(os.environ.get("MODEL_BUK_MK2_OUT") or Path.home() / "Documents" / "Model-Buk-Data-MK2")
DEFAULT_OUT = DEFAULT_OUT_ROOT / "Zaklady MK2.xlsx"
DEFAULT_ROWS = 1000
MIN_ROWS = 50
DEFAULT_FIRST_MONTH = "2026-10"
MONTHS = 12

BETS = "Zakłady"
SUMMARY = "Podsumowanie"
TABLE_NAME = "Zaklady"
AUTHOR = "Model Buk MK2"

TRYBY = ("Realny", "Papier")
RODZAJE = ("Singiel", "Kupon 2", "Kupon 3")
ZRODLA = ("Model", "Własny")
WYNIKI = ("W", "P", "Z")

# number formats (stored in the file in the invariant form; Polish Excel shows "1 234,56 zł")
MONEY = '#,##0.00 "zł"'
ODDS = "0.00"
PCT = "0%"
COUNT = "0"
DATE_FMT = "dd.mm.yyyy"
TIME_FMT = "hh:mm"
MONTH_FMT = "[$-415]mmmm yyyy"

# palette: navy headers for inputs, slate for computed columns, light grey computed cells
NAVY = "1F3864"
SLATE = "44546A"
CALC_FILL = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
NOTE_FILL = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
TABLE_HEAD_FILL = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
GOOD_FONT = Font(color="006100")
GOOD_FILL = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
BAD_FONT = Font(color="9C0006")
BAD_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
MISSING_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")

FONT = Font(name="Arial", size=10)
FONT_CALC = Font(name="Arial", size=10, color="404040")
FONT_BOLD = Font(name="Arial", size=10, bold=True)
FONT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONT_NOTE = Font(name="Arial", size=9, italic=True, color="7F7F7F")
GRID = Side(style="thin", color="D9D9D9")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
ROW_LINE = Border(bottom=GRID)

# ---------------------------------------------------------------------------------------------------------------
# Sheet "Zakłady".  {r} is the row number.  The tax constants are written into each formula on purpose (0.12 and
# 0.88 for the stake tax, 2280 zł limit and 10% for the winnings tax) so every formula can be read and checked alone.
F_LP = '=IF(COUNTA(B{r},C{r},D{r},E{r},F{r},G{r},H{r},I{r},J{r})=0,"",ROW()-1)'
F_TAX12 = '=IF(ISNUMBER(J{r}),J{r}*0.12,"")'
F_NET_STAKE = '=IF(ISNUMBER(J{r}),J{r}*0.88,"")'
F_PAYOUT = ('=IF(OR(M{r}="",NOT(ISNUMBER(J{r}))),"",'
            'IF(M{r}="W",IF(ISNUMBER(I{r}),J{r}*0.88*I{r},""),IF(M{r}="P",0,IF(M{r}="Z",J{r},""))))')
F_WIN_TAX = '=IF(ISNUMBER(N{r}),IF(AND(M{r}="W",N{r}>2280),N{r}*0.1,0),"")'
F_PROFIT = '=IF(ISNUMBER(N{r}),N{r}-O{r}-J{r},"")'
F_BALANCE = '=IF(AND(ISNUMBER(P{r}),U{r}<>""),SUMIFS(P$2:P{r},U$2:U{r},U{r}),"")'
F_MIN_ODDS = '=IF(AND(ISNUMBER(R{r}),R{r}>0),1/(R{r}*0.88),"")'


class Col(NamedTuple):
    key: str
    header: str
    width: float
    formula: str | None  # None = input column typed by the user
    fmt: str | None = None
    wrap: bool = False
    center: bool = False
    note: str | None = None


COLUMNS: tuple[Col, ...] = (
    Col("lp", "Lp", 5, F_LP, COUNT, center=True,
        note="Liczy się samo: numer zakładu, pojawia się po wpisaniu danych w wierszu."),
    Col("data", "Data", 11, None, DATE_FMT, center=True,
        note="Data meczu, np. 14.10.2026. Z niej liczone jest zestawienie miesięczne."),
    Col("godzina", "Godzina", 8, None, TIME_FMT, center=True, note="Godzina rozpoczęcia meczu, np. 21:00."),
    Col("rozgrywki", "Rozgrywki", 14, None),
    Col("mecz", "Mecz", 34, None, wrap=True,
        note="Gospodarz – gość. Kupon: wszystkie mecze w jednej komórce, np. „A – B + C – D”."),
    Col("rodzaj", "Rodzaj", 10, None, center=True, note="Singiel / Kupon 2 / Kupon 3 – wybierz z listy."),
    Col("zaklad", "Zakład", 38, None, wrap=True,
        note="Opis zdarzenia, np. „Strzały łącznie poniżej 25,5”. Kupon: wszystkie zdarzenia w jednej komórce."),
    Col("bukmacher", "Bukmacher", 14, None),
    Col("kurs", "Kurs", 7, None, ODDS,
        note="Kurs u bukmachera (kupon: kurs łączny). Czerwone tło: kurs niższy niż Min. kurs wg modelu."),
    Col("stawka", "Stawka (zł)", 12, None, MONEY,
        note="Kwota postawiona u bukmachera, przed potrąceniem podatku 12%."),
    Col("podatek12", "Podatek 12% (zł)", 11, F_TAX12, MONEY,
        note="Liczy się samo: 12% stawki, pobierane przy zawarciu zakładu (przy zwrocie oddawane)."),
    Col("stawka_netto", "Stawka po podatku (zł)", 12, F_NET_STAKE, MONEY, note="Liczy się samo: 0,88 × stawka."),
    Col("wynik", "Wynik (W/P/Z)", 9, None, center=True,
        note="W = wygrana, P = przegrana, Z = zwrot, puste = nierozliczony. "
             "Żółte tło: data meczu minęła, a wyniku brak."),
    Col("wyplata", "Wypłata (zł)", 13, F_PAYOUT, MONEY,
        note="Liczy się samo: W = 0,88 × stawka × kurs, P = 0, Z = cała stawka. Kwota przed podatkiem od wygranej."),
    Col("podatek_wygr", "Podatek od wygranej (zł)", 12, F_WIN_TAX, MONEY,
        note="Liczy się samo: 10% całej wypłaty, gdy wygrana przekracza 2280 zł; w innym razie 0."),
    Col("zysk", "Zysk (zł)", 13, F_PROFIT, MONEY,
        note="Liczy się samo: wypłata − podatek od wygranej − stawka."),
    Col("saldo", "Saldo (zł)", 13, F_BALANCE, MONEY,
        note="Liczy się samo: suma zysków rozliczonych zakładów od góry tabeli do tego wiersza, "
             "osobno dla trybu Realny i Papier."),
    Col("szansa", "Szansa wg modelu (%)", 10, None, PCT, center=True, note="Prawdopodobieństwo z modelu, np. 47%."),
    Col("min_kurs", "Min. kurs wg modelu", 10, F_MIN_ODDS, ODDS,
        note="Liczy się samo: 1 / (szansa × 0,88) – kurs, przy którym zakład wychodzi na zero po podatku 12%. "
             "Graj tylko przy kursie wyższym."),
    Col("zrodlo", "Źródło", 9, None, center=True, note="Model = typ z modelu, Własny = własny pomysł."),
    Col("tryb", "Tryb", 9, None, center=True,
        note="Realny = prawdziwe pieniądze, Papier = zakład testowy. Podsumowanie liczy je osobno."),
    Col("notatki", "Notatki", 32, None, wrap=True),
)

# The formulas above address columns by letter, so the column order is fixed (A = lp ... V = notatki).
_KEY_ORDER = ("lp data godzina rozgrywki mecz rodzaj zaklad bukmacher kurs stawka podatek12 stawka_netto wynik "
              "wyplata podatek_wygr zysk saldo szansa min_kurs zrodlo tryb notatki").split()
assert [c.key for c in COLUMNS] == _KEY_ORDER, "column order changed - update the formula letters"
LETTER = {c.key: get_column_letter(i) for i, c in enumerate(COLUMNS, start=1)}
LAST_COL = get_column_letter(len(COLUMNS))

SEED_COMMON = dict(data=date(2026, 10, 14), godzina=time(21, 0), rozgrywki="Liga Mistrzów", rodzaj="Singiel",
                   bukmacher="do sprawdzenia", stawka=10, zrodlo="Model", tryb="Papier")
SEED_ROWS: tuple[dict, ...] = (
    dict(SEED_COMMON, mecz="Aston Villa – Fenerbahçe", zaklad="Celne strzały łącznie powyżej 10,5",
         kurs=2.67, szansa=0.47),
    dict(SEED_COMMON, mecz="Bodø/Glimt – Borussia Dortmund", zaklad="Strzały łącznie poniżej 25,5",
         kurs=2.50, szansa=0.49),
    dict(SEED_COMMON, mecz="Szachtar Donieck – AEK Ateny", zaklad="Rożne łącznie poniżej 7,5",
         kurs=2.95, szansa=0.42),
    dict(SEED_COMMON, mecz="Manchester City – Paris Saint-Germain", zaklad="Strzały łącznie poniżej 26,5",
         kurs=2.84, szansa=0.44),
    dict(SEED_COMMON, mecz="Slovan Bratysława – VfB Stuttgart", zaklad="Strzały łącznie powyżej 30,5",
         kurs=2.90, szansa=0.54),
    dict(SEED_COMMON, rodzaj="Kupon 2", mecz="Aston Villa – Fenerbahçe + Roma – Real Madryt",
         zaklad="Aston Villa powyżej 13,5 strzałów + Real Madryt powyżej 3,5 celnych strzałów",
         kurs=None, szansa=0.60, notatki="graj tylko przy kursie ≥ 1,90"),
)

HOW_TO = (
    "Każdy zakład wpisz w arkuszu „Zakłady” w pierwszym wolnym wierszu: wypełniasz tylko białe kolumny, a szare "
    "(podatek, wypłata, zysk, saldo, minimalny kurs) liczą się same; żółte pole oznacza brakującą daną. "
    "Po meczu wpisz w kolumnie „Wynik” literę W (wygrana), P (przegrana) lub Z (zwrot) – puste pole oznacza zakład "
    "jeszcze nierozliczony. W kolumnie „Tryb” wybierz „Realny” dla zakładów za prawdziwe pieniądze albo „Papier” "
    "dla zakładów testowych – poniżej oba tryby są liczone osobno. Podatek 12% od stawki i 10% od wygranej "
    "powyżej 2280 zł są uwzględniane automatycznie."
)
FOOTNOTE = (
    "Zasady liczenia: przy zawarciu zakładu bukmacher pobiera 12% stawki (gra 0,88 × stawka); wygrana wypłaca "
    "0,88 × stawka × kurs, zwrot oddaje całą stawkę; od wygranej powyżej 2280 zł pobierany jest 10% podatku od całej "
    "kwoty. Zestawienia liczą tylko wiersze z wybranym Trybem; Suma stawek i ROI obejmują zakłady rozliczone "
    "(także zwroty), Trafność = W / (W + P)."
)


def excel_serial(day: date) -> int:
    return (day - date(1899, 12, 30)).days


def build_bets_sheet(ws: Worksheet, n_rows: int, seed: bool) -> None:
    last = n_rows + 1
    for idx, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        cell = ws.cell(row=1, column=idx, value=col.header)
        cell.font = FONT_HEADER
        cell.fill = PatternFill("solid", start_color=SLATE if col.formula else NAVY,
                                end_color=SLATE if col.formula else NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        if col.note:
            cell.comment = Comment(col.note, AUTHOR, height=90, width=230)
        ws.column_dimensions[letter].width = col.width
    ws.row_dimensions[1].height = 42

    aligns = [Alignment(horizontal="center" if c.center else None, vertical="center", wrap_text=c.wrap or None)
              for c in COLUMNS]
    seeds = SEED_ROWS if seed else ()
    for r in range(2, last + 1):
        values = seeds[r - 2] if r - 2 < len(seeds) else {}
        for idx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx)
            if col.formula:
                cell.value = col.formula.format(r=r)
                cell.font = FONT_CALC
                cell.fill = CALC_FILL
            else:
                cell.value = values.get(col.key)
                cell.font = FONT
            if col.fmt:
                cell.number_format = col.fmt
            cell.alignment = aligns[idx - 1]
            cell.border = BORDER

    table = Table(displayName=TABLE_NAME, ref=f"A1:{LAST_COL}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False, showLastColumn=False,
                                          showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)  # a table with a header row gets its own filter buttons
    ws.freeze_panes = "F2"  # header row plus Lp..Mecz stay visible while scrolling

    add_validations(ws, last)
    add_bets_formatting(ws, last)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"


def add_validations(ws: Worksheet, last: int) -> None:
    def add(col_key: str, dv: DataValidation, title: str, error: str, prompt: str | None = None) -> None:
        dv.allow_blank = True
        dv.showErrorMessage = True
        dv.errorStyle = "stop"
        dv.errorTitle = title
        dv.error = error
        if prompt:
            dv.showInputMessage = True
            dv.promptTitle = title
            dv.prompt = prompt
        letter = LETTER[col_key]
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)

    def choice(items: tuple[str, ...]) -> DataValidation:
        return DataValidation(type="list", formula1='"' + ",".join(items) + '"')

    add("rodzaj", choice(RODZAJE), "Rodzaj", "Wybierz z listy: Singiel, Kupon 2 albo Kupon 3.")
    add("wynik", choice(WYNIKI), "Wynik",
        "Wpisz W (wygrana), P (przegrana) lub Z (zwrot) albo zostaw puste.",
        prompt="W = wygrana, P = przegrana, Z = zwrot. Zostaw puste, dopóki zakład nie jest rozliczony.")
    add("zrodlo", choice(ZRODLA), "Źródło", "Wybierz z listy: Model albo Własny.")
    add("tryb", choice(TRYBY), "Tryb", "Wybierz z listy: Realny (prawdziwe pieniądze) albo Papier (test).")
    add("data", DataValidation(type="date", operator="between", formula1=str(excel_serial(date(2020, 1, 1))),
                               formula2=str(excel_serial(date(2100, 12, 31)))),
        "Data", "Wpisz datę meczu, np. 14.10.2026.")
    add("godzina", DataValidation(type="time", operator="between", formula1="0", formula2="0.999988425925926"),
        "Godzina", "Wpisz godzinę, np. 21:00.")
    add("kurs", DataValidation(type="decimal", operator="between", formula1="1.01", formula2="1000"),
        "Kurs", "Wpisz kurs jako liczbę z przecinkiem, np. 2,67 (od 1,01 do 1000).")
    add("stawka", DataValidation(type="decimal", operator="between", formula1="0.01", formula2="1000000"),
        "Stawka", "Wpisz stawkę w zł jako liczbę, np. 10 albo 12,50.")
    add("szansa", DataValidation(type="decimal", operator="between", formula1="0", formula2="1"),
        "Szansa wg modelu", "Wpisz szansę w procentach, np. 47% (od 0% do 100%).",
        prompt="Szansa z modelu w procentach, np. 47%.")


def add_rule(ws: Worksheet, sqref: str, formula: str, font: Font | None = None,
             fill: PatternFill | None = None) -> None:
    ws.conditional_formatting.add(sqref, FormulaRule(formula=[formula], font=font, fill=fill))


def add_bets_formatting(ws: Worksheet, last: int) -> None:
    # Zysk and Saldo: green when positive, red when negative (blank cells hold "" and are skipped by ISNUMBER)
    add_rule(ws, f"P2:Q{last}", "AND(ISNUMBER(P2),P2>0)", GOOD_FONT, GOOD_FILL)
    add_rule(ws, f"P2:Q{last}", "AND(ISNUMBER(P2),P2<0)", BAD_FONT, BAD_FILL)
    # result letters
    add_rule(ws, f"M2:M{last}", 'M2="W"', Font(bold=True, color="006100"))
    add_rule(ws, f"M2:M{last}", 'M2="P"', Font(bold=True, color="9C0006"))
    add_rule(ws, f"M2:M{last}", 'M2="Z"', Font(bold=True, color="7F7F7F"))
    # match date passed but no result yet
    add_rule(ws, f"M2:M{last}", 'AND($A2<>"",M2="",ISNUMBER($B2),$B2<TODAY())', fill=MISSING_FILL)
    # odds below the model's break-even odds
    add_rule(ws, f"I2:I{last}", "AND(ISNUMBER(I2),ISNUMBER($S2),I2<$S2)", BAD_FONT, BAD_FILL)
    # required inputs missing in a row that is in use
    for key in ("data", "rodzaj", "kurs", "stawka", "zrodlo", "tryb"):
        col = LETTER[key]
        add_rule(ws, f"{col}2:{col}{last}", f'AND($A2<>"",{col}2="")', fill=MISSING_FILL)


# ---------------------------------------------------------------------------------------------------------------
# Sheet "Podsumowanie"

def zr(col: str) -> str:
    """A whole column of the bets sheet, e.g. 'Zakłady'!$J:$J (rows added below the table still count)."""
    return f"'{BETS}'!${col}:${col}"


def crit(col: str, criterion: str) -> str:
    return f"{zr(col)},{criterion}"


def q(text: str) -> str:
    return f'"{text}"'


SETTLED = crit("N", q(">=0"))  # Wypłata is a number (>= 0) exactly when the bet is settled

MAIN_METRICS = (
    # key, label, number format, explanation shown next to the value
    ("n", "Liczba zakładów", COUNT, "wszystkie wpisane zakłady"),
    ("open", "W grze (nierozliczone)", COUNT, "zakłady bez wyniku"),
    ("settled", "Rozliczone", COUNT, "W + P + Z"),
    ("w", "Wygrane (W)", COUNT, None),
    ("p", "Przegrane (P)", COUNT, None),
    ("z", "Zwroty (Z)", COUNT, None),
    ("hit", "Trafność", PCT, "W / (W + P), zwroty pomijane"),
    ("stakes", "Suma stawek (zł)", MONEY, "stawki rozliczonych zakładów (z podatkiem 12%)"),
    ("inplay", "Stawki w grze (zł)", MONEY, "stawki zakładów bez wyniku"),
    ("payouts", "Suma wypłat (zł)", MONEY, "przed podatkiem od wygranej; zwrot = cała stawka"),
    ("tax12", "Podatek 12% (zł)", MONEY, "od stawek wygranych i przegranych (przy zwrocie oddawany)"),
    ("taxwin", "Podatek od wygranych (zł)", MONEY, "10% całej wypłaty, gdy przekracza 2280 zł"),
    ("taxall", "Podatek łącznie (zł)", MONEY, "podatek 12% + podatek od wygranych"),
    ("profit", "Zysk (zł)", MONEY, "wypłaty − podatek od wygranych − stawki"),
    ("roi", "ROI", PCT, "zysk / suma stawek"),
    ("balance", "Saldo bieżące (zł)", MONEY, "ostatnie saldo z arkusza Zakłady (zysk narastająco)"),
)
HIGHLIGHT = ("profit", "roi", "balance")


def main_formulas(tryb: str, cell: dict[str, str]) -> dict[str, str]:
    t = crit("U", q(tryb))
    won, lost, void = crit("M", q("W")), crit("M", q("P")), crit("M", q("Z"))
    return {
        "n": f"=COUNTIFS({t},{crit('A', q('>0'))})",
        "open": f"={cell['n']}-{cell['settled']}",
        "settled": f"=COUNTIFS({t},{SETTLED})",
        "w": f"=COUNTIFS({t},{won},{SETTLED})",
        "p": f"=COUNTIFS({t},{lost},{SETTLED})",
        "z": f"=COUNTIFS({t},{void},{SETTLED})",
        "hit": f'=IF({cell["w"]}+{cell["p"]}=0,"",{cell["w"]}/({cell["w"]}+{cell["p"]}))',
        "stakes": f"=SUMIFS({zr('J')},{t},{SETTLED})",
        "inplay": f"=SUMIFS({zr('J')},{t})-{cell['stakes']}",
        "payouts": f"=SUMIFS({zr('N')},{t})",
        "tax12": f"=SUMIFS({zr('K')},{t},{won},{SETTLED})+SUMIFS({zr('K')},{t},{lost},{SETTLED})",
        "taxwin": f"=SUMIFS({zr('O')},{t})",
        "taxall": f"={cell['tax12']}+{cell['taxwin']}",
        "profit": f"=SUMIFS({zr('P')},{t})",
        "roi": f'=IF({cell["stakes"]}=0,"",{cell["profit"]}/{cell["stakes"]})',
        "balance": f"={cell['profit']}",
    }


BREAKDOWN_HEADERS = ("Zakłady", "Rozliczone", "Trafność", "Stawki (zł)", "Zysk (zł)", "ROI")
BREAKDOWN_FORMATS = (COUNT, COUNT, PCT, MONEY, MONEY, PCT)


def breakdown_formulas(tryb: str, extra: str, stakes_cell: str, profit_cell: str) -> list[str]:
    base = f"{crit('U', q(tryb))},{extra}"
    wins = f"COUNTIFS({base},{crit('M', q('W'))},{SETTLED})"
    losses = f"COUNTIFS({base},{crit('M', q('P'))},{SETTLED})"
    return [
        f"=COUNTIFS({base},{crit('A', q('>0'))})",
        f"=COUNTIFS({base},{SETTLED})",
        f'=IFERROR({wins}/({wins}+{losses}),"")',
        f"=SUMIFS({zr('J')},{base},{SETTLED})",
        f"=SUMIFS({zr('P')},{base})",
        f'=IF({stakes_cell}=0,"",{profit_cell}/{stakes_cell})',
    ]


class SummaryWriter:
    """Writes one side (Realny or Papier) of the summary: label column `lc` plus six metric columns."""

    def __init__(self, ws: Worksheet, tryb: str, lc: int, colour: str) -> None:
        self.ws, self.tryb, self.lc, self.colour = ws, tryb, lc, colour
        self.L = get_column_letter(lc)
        self.span = f"{self.L}{{r}}:{get_column_letter(lc + 6)}{{r}}"
        self.highlight: list[str] = []  # ranges coloured green/red

    def section(self, row: int, title: str) -> None:
        ws = self.ws
        ws.merge_cells(self.span.format(r=row))
        cell = ws.cell(row=row, column=self.lc, value=title)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=self.colour, end_color=self.colour)
        cell.alignment = Alignment(vertical="center", indent=1)

    def subheader(self, row: int, title: str) -> None:
        for c in range(self.lc, self.lc + 7):
            self.ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=self.colour))
        cell = self.ws.cell(row=row, column=self.lc, value=title)
        cell.font = Font(name="Arial", size=10, bold=True, color=self.colour)

    def main_block(self, top: int) -> int:
        ws, vc = self.ws, self.lc + 1
        V = get_column_letter(vc)
        rows = {key: top + i for i, (key, *_rest) in enumerate(MAIN_METRICS)}
        cells = {key: f"{V}{r}" for key, r in rows.items()}
        formulas = main_formulas(self.tryb, cells)
        for key, label, fmt, note in MAIN_METRICS:
            r = rows[key]
            strong = key in HIGHLIGHT
            lab = ws.cell(row=r, column=self.lc, value=label)
            lab.font = FONT_BOLD if strong else FONT
            val = ws.cell(row=r, column=vc, value=formulas[key])
            val.font = FONT_BOLD if strong else FONT
            val.number_format = fmt
            val.alignment = Alignment(horizontal="right")
            for c in range(self.lc, self.lc + 7):
                ws.cell(row=r, column=c).border = ROW_LINE
            if note:
                ws.cell(row=r, column=vc + 1, value=note).font = FONT_NOTE
        self.highlight.append(f"{cells[HIGHLIGHT[0]]}:{cells[HIGHLIGHT[-1]]}")  # Zysk, ROI, Saldo (consecutive)
        return top + len(MAIN_METRICS)

    def table_header(self, row: int, first: str) -> None:
        for i, text in enumerate((first, *BREAKDOWN_HEADERS)):
            cell = self.ws.cell(row=row, column=self.lc + i, value=text)
            cell.font = FONT_BOLD
            cell.fill = TABLE_HEAD_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True)

    def breakdown_row(self, r: int, label, label_fmt: str | None, extra: str) -> None:
        ws = self.ws
        lab = ws.cell(row=r, column=self.lc, value=label)
        lab.font, lab.border = FONT, BORDER
        lab.alignment = Alignment(horizontal="left")
        if label_fmt:
            lab.number_format = label_fmt
        stakes = f"{get_column_letter(self.lc + 4)}{r}"
        profit = f"{get_column_letter(self.lc + 5)}{r}"
        for i, (formula, fmt) in enumerate(zip(breakdown_formulas(self.tryb, extra, stakes, profit),
                                               BREAKDOWN_FORMATS), start=1):
            cell = ws.cell(row=r, column=self.lc + i, value=formula)
            cell.font, cell.border, cell.number_format = FONT, BORDER, fmt

    def highlight_block(self, first_row: int, last_row: int) -> None:
        """Zysk and ROI columns of a breakdown table get the green/red rule."""
        self.highlight.append(f"{get_column_letter(self.lc + 5)}{first_row}:{get_column_letter(self.lc + 6)}{last_row}")

    def category_table(self, top: int, title: str, first: str, key_col: str, labels: tuple[str, ...]) -> int:
        self.subheader(top, title)
        self.table_header(top + 1, first)
        for i, label in enumerate(labels):
            r = top + 2 + i
            self.breakdown_row(r, label, None, crit(key_col, f"${self.L}{r}"))
        self.highlight_block(top + 2, top + 1 + len(labels))
        return top + 2 + len(labels)

    def month_table(self, top: int, first_month: date, mirror_of: str | None) -> int:
        self.subheader(top, "Wg miesiąca (data meczu)")
        self.table_header(top + 1, "Miesiąc")
        for i in range(MONTHS):
            r = top + 2 + i
            if mirror_of:
                label = f"=${mirror_of}{r}"
            elif i == 0:
                label = first_month
            else:
                label = f"=DATE(YEAR({self.L}{r - 1}),MONTH({self.L}{r - 1})+1,1)"
            start = f"${self.L}{r}"
            next_month = f"DATE(YEAR({start}),MONTH({start})+1,1)"
            extra = f"{crit('B', q('>=') + '&' + start)},{crit('B', q('<') + '&' + next_month)}"
            self.breakdown_row(r, label, MONTH_FMT, extra)
        if not mirror_of:
            self.ws.cell(row=top + 2, column=self.lc).comment = Comment(
                "Pierwszy miesiąc zestawienia. Wpisz inną datę (1. dzień miesiąca), aby przesunąć okres; "
                "kolejne miesiące i strona Papier przeliczą się same.", AUTHOR, height=80, width=230)
        self.highlight_block(top + 2, top + 1 + MONTHS)
        return top + 2 + MONTHS


def build_summary_sheet(ws: Worksheet, first_month: date) -> None:
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 26, "C": 14, "D": 11, "E": 10, "F": 13, "G": 13, "H": 9, "I": 3,
              "J": 26, "K": 14, "L": 11, "M": 10, "N": 13, "O": 13, "P": 9}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    ws["B1"] = "Podsumowanie zakładów"
    ws["B1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 24
    ws["B2"] = "Jak używać"
    ws["B2"].font = FONT_BOLD
    ws.merge_cells("B3:P3")
    ws["B3"] = HOW_TO
    ws["B3"].font = FONT
    ws["B3"].fill = NOTE_FILL
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 54

    sides = (SummaryWriter(ws, "Realny", 2, NAVY), SummaryWriter(ws, "Papier", 10, SLATE))
    titles = ("REALNY – zakłady za prawdziwe pieniądze", "PAPIER – zakłady testowe (bez pieniędzy)")
    end = 0
    for side, title in zip(sides, titles):
        side.section(5, title)
        side.subheader(6, "Wyniki ogółem")
        row = side.main_block(7) + 1
        row = side.category_table(row, "Wg rodzaju zakładu", "Rodzaj", LETTER["rodzaj"], RODZAJE) + 1
        row = side.category_table(row, "Wg źródła typu", "Źródło", LETTER["zrodlo"], ZRODLA) + 1
        row = side.month_table(row, first_month, None if side.tryb == TRYBY[0] else sides[0].L)
        end = row
        for ref in side.highlight:
            first = ref.split(":")[0]
            add_rule(ws, ref, f"AND(ISNUMBER({first}),{first}>0)", GOOD_FONT, GOOD_FILL)
            add_rule(ws, ref, f"AND(ISNUMBER({first}),{first}<0)", BAD_FONT, BAD_FILL)
    ws.row_dimensions[5].height = 20

    foot = end + 1
    ws.merge_cells(f"B{foot}:P{foot}")
    ws[f"B{foot}"] = FOOTNOTE
    ws[f"B{foot}"].font = FONT_NOTE
    ws[f"B{foot}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[foot].height = 30

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def build_workbook(n_rows: int = DEFAULT_ROWS, first_month: date = date(2026, 10, 1), seed: bool = True) -> Workbook:
    wb = Workbook()
    bets = wb.active
    bets.title = BETS
    build_bets_sheet(bets, n_rows, seed)
    build_summary_sheet(wb.create_sheet(SUMMARY), first_month)
    wb.active = 0
    wb.properties.title = "Zakłady MK2"
    wb.properties.creator = AUTHOR
    return wb


def write_new(path: Path, data: bytes, force: bool) -> Path | None:
    """Write `data` to `path`.  Without force an existing file is left alone; with force it is backed up first."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if not force:
        with open(path, "xb") as fh:  # "x" fails if the file appeared meanwhile: never overwrite user entries
            fh.write(data)
        return None
    backup = None
    if path.exists():
        backup = path.with_name(f"{path.stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        shutil.copy2(path, backup)
    tmp = path.with_name(path.name + ".tmp")
    try:
        tmp.write_bytes(data)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
    return backup


def parse_month(text: str) -> date:
    try:
        stamp = datetime.strptime(text, "%Y-%m")
    except ValueError as exc:
        raise argparse.ArgumentTypeError(f"expected YYYY-MM, got {text!r}") from exc
    return date(stamp.year, stamp.month, 1)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help=f"workbook path (default: {DEFAULT_OUT})")
    parser.add_argument("--rows", type=int, default=DEFAULT_ROWS,
                        help=f"rows with pre-filled formulas (default {DEFAULT_ROWS}, minimum {MIN_ROWS})")
    parser.add_argument("--first-month", type=parse_month, default=parse_month(DEFAULT_FIRST_MONTH),
                        help=f"first month of the monthly table, YYYY-MM (default {DEFAULT_FIRST_MONTH})")
    parser.add_argument("--no-seed", action="store_true", help="leave out the six paper example bets")
    parser.add_argument("--force", action="store_true",
                        help="overwrite an existing workbook (the old file is copied to a .backup-<time>.xlsx first)")
    args = parser.parse_args(argv)
    if args.rows < MIN_ROWS:
        parser.error(f"--rows must be at least {MIN_ROWS}")

    out: Path = args.out
    if out.exists() and not args.force:
        print(f"Plik już istnieje: {out}\nNie nadpisuję go, żeby nie stracić wpisanych zakładów. "
              "Aby zbudować go od nowa, uruchom z --force (stary plik zostanie zachowany jako kopia .backup).")
        return 2

    wb = build_workbook(args.rows, args.first_month, seed=not args.no_seed)
    buf = io.BytesIO()
    wb.save(buf)
    try:
        backup = write_new(out, buf.getvalue(), args.force)
    except FileExistsError:
        print(f"Plik już istnieje: {out} - nie nadpisuję (użyj --force).")
        return 2
    except PermissionError as exc:
        print(f"Nie mogę zapisać {out}: {exc}. Jeśli plik jest otwarty w Excelu, zamknij go i spróbuj ponownie.")
        return 1
    seeds = 0 if args.no_seed else len(SEED_ROWS)
    print(f"Zapisano: {out}  (wiersze z formułami: 2-{args.rows + 1}, przykładowe zakłady: {seeds}, "
          f"miesiące od {args.first_month:%Y-%m})")
    if backup:
        print(f"Poprzednia wersja zachowana jako: {backup}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
